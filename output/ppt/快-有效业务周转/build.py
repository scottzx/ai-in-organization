#!/usr/bin/env python3
"""Build an editable Chinese presentation, PDF preview, and speaker manuscript."""
from pathlib import Path
from collections import Counter
from urllib.parse import quote
import argparse
import hashlib
import json
import re
import subprocess

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
STEM = '快-让企业更快完成有效业务循环'
W, H = 13.333333, 7.5
C = {'navy':'12213A','ink':'1C2D48','blue':'2869ED','cyan':'5CB8FF',
     'bg':'F4F7FC','white':'FFFFFF','muted':'586C87','line':'D8E3F2',
     'pale':'EAF1FF','teal':'138D91','mint':'E5F5F3','amber':'B17A27',
     'sand':'FFF3DF','red':'AD5360','rose':'FAEBEF'}
FONT = 'Noto Sans CJK SC'
EXPECTED = []
SHAPES = []
DATA = json.loads((HERE/'content.json').read_text(encoding='utf-8'))
PRS = Presentation()
PRS.slide_width, PRS.slide_height = Inches(W), Inches(H)
PRS.core_properties.title = DATA['title']
PRS.core_properties.subject = DATA['subtitle']
PRS.core_properties.author = DATA['author']
PRS.core_properties.keywords = '快,有效业务周转,MTL,LTC,智力飞轮,决策,试错,复制,切换'


def color(name):
    return RGBColor.from_string(C.get(name, name))


def rect(s, x, y, w, h, fill='white', line=None, radius=False):
    shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
    if line:shape.line.color.rgb=color(line);shape.line.width=Pt(.8)
    else:shape.line.fill.background()
    if radius:shape.adjustments[0]=.12
    return shape


def txt(s, x, y, w, h, value, size=20, ink='ink', bold=False, align='left', margin=.02, url=None):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    box.name=f'Text {len(EXPECTED)+1}'
    tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin)
    tf.margin_top=tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    lines=str(value).split('\n')
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.14
        run=p.add_run();run.text=line
        run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color(ink)
        ea=OxmlElement('a:ea');ea.set('typeface',FONT);run._r.get_or_add_rPr().append(ea)
        if url:run.hyperlink.address=url
    EXPECTED.append({'slide':len(PRS.slides),'text':str(value),'shape':box.name})
    SHAPES.append({'slide':len(PRS.slides),'shape':box.name,'x':x,'y':y,'w':w,'h':h})
    return box


def line(s,x1,y1,x2,y2,ink='line',width=2):
    shape=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    shape.line.color.rgb=color(ink);shape.line.width=Pt(width)
    return shape


def arrow(s,x,y,w=.24,h=.22,ink='blue'):
    return rect_shape(s,MSO_SHAPE.CHEVRON,x,y,w,h,ink)


def rect_shape(s,kind,x,y,w,h,fill):
    shape=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=color(fill);shape.line.fill.background()
    return shape


def repository_source_url(key):
    source=DATA['sources'][key]
    url='https://github.com/scottzx/ai-in-organization/blob/'+DATA['source_revision']+'/'+quote(source['path'])
    return url+('#'+source['anchor'] if source.get('anchor') else '')


def source_url(key):
    source=DATA['sources'][key]
    return source.get('url') or repository_source_url(key)


def source_bytes(source):
    """Read the presentation's fixed source revision, independent of the worktree."""
    return subprocess.run(['git','show',DATA['source_revision']+':'+source['path']],
                          cwd=ROOT,check=True,capture_output=True).stdout


def base(d, dark=False):
    s=PRS.slides.add_slide(PRS.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=color('navy' if dark else 'bg')
    ink='white' if dark else 'ink'
    footer_ink='B6C9E4' if dark else 'muted'
    if not dark:
        txt(s,.64,.35,10,.25,'一芥智能研究院  /  企业数智化转型研究报告',10.5,'muted')
        rect(s,.66,.84,.07,.62,'blue')
        txt(s,.89,.77,11.85,.73,d['title'],30,ink,True)
        txt(s,.9,1.53,11.73,.5,d['subtitle'],15.5,'muted')
    line(s,.65,6.98,12.69,6.98,'36506E' if dark else 'line',.7)
    key=next((r for r in d['refs'] if r in DATA['sources'] and r[0].isupper()),d['refs'][0])
    txt(s,.65,7.08,11.55,.23,d['footer'],9.0,footer_ink,url=None if dark else source_url(key))
    txt(s,12.23,7.05,.46,.27,f'{d["number"]:02d}',10.5,footer_ink,align='right')
    notes=s.notes_slide.notes_text_frame
    note=d['notes']+'\n\n【依据与定位】\n'
    for ref in d['refs']:
        entry=DATA['sources'][ref]
        note+=entry['label']+'\n'+entry['path']
        if entry.get('anchor'):note+='#'+entry['anchor']
        if entry.get('stage'):note+='｜阶段：'+entry['stage']
        note+='\n'+source_url(ref)+'\n'
        if entry.get('locators'):note+='；'.join(entry['locators'])+'\n'
    notes.text=note
    return s


def message(s,text,y=6.2):
    rect(s,.67,y,12.0,.56,'pale',radius=True)
    txt(s,.89,y+.05,11.56,.46,text,17,'blue',True,align='center')


def flow(s,labels,y,x=.75,total=11.83,h=.96,size=20,fill='white'):
    n=len(labels);gap=.34;w=(total-gap*(n-1))/n
    for i,label in enumerate(labels):
        xx=x+i*(w+gap)
        rect(s,xx,y,w,h,fill,'line',True)
        txt(s,xx+.07,y+.12,w-.14,h-.24,label,size,'ink',True,'center')
        if i<n-1:arrow(s,xx+w+.075,y+h/2-.09,.18,.18)


def render(d):
    k=d['kind']
    s=base(d,k in ('cover','closing'))
    if k=='cover':
        txt(s,.82,.73,6.7,.62,'解决方案，就是',30,'white',True)
        txt(s,.6,1.39,6.15,3.85,'快',200,'cyan',True)
        txt(s,.89,5.3,6.25,.62,d['subtitle'],25,'white',True)
        txt(s,.92,6.08,5.8,.34,'20分钟主题分享  ·  2026.10.14',12,'B6C9E4')
        for i,label in enumerate(d['speeds']):
            yy=1.67+i*1.12
            rect(s,8.13,yy,3.9,.89,'1D3556',radius=True)
            txt(s,8.32,yy+.16,.5,.53,f'{i+1:02d}',17,'cyan',True)
            txt(s,9.01,yy+.12,2.65,.61,label,26,'white',True)
            if i<3:line(s,10.08,yy+.91,10.08,yy+1.1,'557395',2)
    elif k=='cycle':
        flow(s,d['nodes'],2.62,h=1.25,size=20)
        line(s,11.51,4.0,11.51,4.86,'blue',2)
        line(s,11.51,4.86,1.75,4.86,'blue',2)
        line(s,1.75,4.86,1.75,4.0,'blue',2)
        txt(s,3.63,4.24,6.1,.48,'反馈改变下一轮的输入与方法',20,'blue',True,'center')
        message(s,d['message'])
    elif k=='report_path':
        for i,(number,label,detail) in enumerate(d['steps']):
            yy=2.16+i*.76
            rect(s,.89,yy,.59,.52,'blue',radius=True)
            txt(s,.94,yy+.03,.48,.44,number,17,'white',True,'center')
            txt(s,1.73,yy+.01,3.17,.5,label,23,'ink',True)
            txt(s,5.05,yy+.06,4.42,.39,detail,18,'muted')
        line(s,9.43,2.48,10.03,2.48,'blue',2)
        line(s,10.03,2.48,10.03,5.58,'blue',2)
        line(s,9.43,5.58,10.03,5.58,'blue',2)
        txt(s,10.45,3.17,1.63,2.2,'快',85,'blue',True,'center')
        txt(s,10.14,5.37,2.37,.6,'有效业务周转',17,'blue',True,'center')
    elif k=='bottleneck':
        labels=d['steps'];x=.81;gap=.25;ww=[1.91,3.16,3.11,2.75]
        for i,(label,w) in enumerate(zip(labels,ww)):
            rect(s,x,2.39,w,1.03,'pale' if i==0 else ('sand' if i in (1,2) else 'white'),radius=True)
            txt(s,x+.08,2.6,w-.16,.57,label,21,'blue' if i==0 else 'ink',True,'center')
            x+=w+gap
        txt(s,.9,3.87,2.8,1.18,d['stat'],60,'blue',True)
        txt(s,4.15,4.01,7.82,.5,d['stat_label'],25,'ink',True)
        txt(s,4.16,4.65,7.78,.55,'泳装商拍：生成之后，仍有合格交付的工作',17,'muted')
        message(s,d['lesson'])
    elif k=='mtl_ltc':
        txt(s,.81,2.16,1.3,.4,'MTL',24,'blue',True)
        txt(s,2.03,2.17,9.8,.34,'市场与客户开发：形成销售接得住的线索',17,'muted')
        flow(s,d['mtl'],2.86,h=.9,size=19)
        txt(s,.81,4.1,1.3,.4,'LTC',24,'teal',True)
        txt(s,2.03,4.11,9.8,.34,'销售、方案、交付与财务：履约、验收、回款',17,'muted')
        flow(s,d['ltc'],4.82,h=.9,size=19)
        txt(s,7.2,6.12,5.08,.5,'沿同一客户对象接续，反馈回到前端',17,'blue',True,'right')
    elif k=='four_speeds':
        gap=.24;w=2.85
        for i,(label,distance,detail) in enumerate(d['cards']):
            distance=distance.replace(' → ', '\n→ ')
            x=.67+i*(w+gap)
            rect(s,x,2.42,w,3.56,'white','line',True)
            txt(s,x+.25,2.63,w-.5,.47,f'{i+1:02d}',17,'blue',True)
            txt(s,x+.25,3.23,w-.5,.59,label,30,'blue',True)
            txt(s,x+.22,4.17,w-.44,.85,distance,18.2,'ink',True,'center')
            txt(s,x+.21,5.12,w-.42,.55,detail,15.2,'muted',False,'center')
    elif k=='decision':
        for i,(role,detail) in enumerate(d['roles']):
            x=.86+i*4.19
            rect(s,x,2.34,3.39,1.66,'white','line',True)
            txt(s,x+.18,2.59,3.03,.44,role,25,'blue',True,'center')
            txt(s,x+.13,3.23,3.13,.43,detail,18,'ink',False,'center')
            if i<2:arrow(s,x+3.6,3.02,.28,.24)
        txt(s,.92,4.4,11.56,.5,d['case_title'],22,'ink',True)
        txt(s,.92,5.11,11.52,.49,d['case_body'],20,'muted')
        message(s,d['message'])
    elif k=='trial_case':
        flow(s,d['steps'],2.29,h=.9,size=21)
        for i,(label,result) in enumerate(d['results']):
            x=1.61+i*5.34
            rect(s,x,3.67,4.76,1.94,'white','line',True)
            txt(s,x+.22,3.91,4.32,.45,label,21,'muted',False,'center')
            txt(s,x+.22,4.62,4.32,.7,result+(' ↓' if i==0 else ' ↑'),38,'blue' if i==0 else 'teal',True,'center')
        message(s,d['message'])
    elif k=='flow_case':
        txt(s,.83,2.22,5.08,.9,d['before'],43,'muted',True,'center')
        arrow(s,6.12,2.48,.8,.43)
        txt(s,7.19,2.22,5.09,.9,d['after'],54,'blue',True,'center')
        txt(s,.93,3.26,11.55,.42,'确定合作 → 收样 → 创作 → 提交的平均周期',18,'muted',False,'center')
        flow(s,d['steps'],4.22,h=.92,size=21)
        message(s,d['message'])
    elif k=='copy':
        flow(s,d['steps'],2.61,h=1.42,size=20)
        txt(s,.9,4.68,11.55,.74,d['case_title'],20.3,'ink',True,'center')
        message(s,d['message'])
    elif k=='switch':
        for i,(heading,body) in enumerate(d['choices']):
            x=.79+i*4.21
            rect(s,x,2.3,3.54,1.76,'white','line',True)
            txt(s,x+.16,2.59,3.22,.45,heading,24,'blue',True,'center')
            txt(s,x+.16,3.22,3.22,.58,body,17,'muted',False,'center')
        txt(s,.89,4.5,11.59,.65,d['case_title'],21.5,'ink',True,'center')
        rect(s,.87,5.56,11.59,.69,'pale',radius=True)
        txt(s,1.02,5.64,11.29,.53,d['case_body'],17.5,'blue',True,'center')
    elif k=='four_flows':
        # Lines are undirected collaboration/feedback links, not a serial workflow.
        center=(6.67,4.34)
        positions=[(.76,2.22),(8.49,2.22),(.76,4.74),(8.49,4.74)]
        for x,y in positions:line(s,x+1.98,y+.84,*center,'line',2)
        rect_shape(s,MSO_SHAPE.OVAL,5.19,3.29,2.96,2.12,'blue')
        txt(s,5.33,3.88,2.67,.7,d['center'],27,'white',True,'center')
        for (x,y),(code,name,result) in zip(positions,d['flows']):
            rect(s,x,y,4.09,1.84,'white','line',True)
            txt(s,x+.18,y+.16,.96,.42,code,22,'blue',True)
            txt(s,x+1.31,y+.17,2.57,.4,name,19,'ink',True)
            txt(s,x+.18,y+.97,3.7,.57,result,18,'muted',False,'center')
        txt(s,5.17,5.94,3.1,.53,'共同责任 · 交接 · 反馈',15.5,'blue',True,'center')
    elif k=='network':
        stops_x=[6.03,8.11,10.36]
        for i,(actor,role) in enumerate(d['lanes']):
            yy=2.43+i*1.09
            rect(s,.82,yy,2.26,.7,'white','line',True)
            txt(s,.94,yy+.11,2.03,.46,actor,22,'blue',True,'center')
            txt(s,3.38,yy+.14,3.01,.43,role,15.5,'muted')
            line(s,6.43,yy+.38,12.19,yy+.38,'blue',2)
            for xx in stops_x:rect_shape(s,MSO_SHAPE.OVAL,xx+.71,yy+.24,.25,.25,'blue')
        for xx,label in zip(stops_x,d['stops']):
            txt(s,xx+.04,5.59,2.02,.43,label,16.5,'ink',True,'center')
        message(s,d['message'])
    elif k=='flywheel':
        positions=[(.76,2.38),(5.33,2.26),(9.9,2.38),(8.14,5.17),(2.31,5.17)]
        centers=[(x+1.32,y+.45) for x,y in positions]
        for a,b in zip(centers,centers[1:]+centers[:1]):line(s,*a,*b,'blue',2.2)
        for (x,y),label in zip(positions,d['nodes']):
            rect(s,x,y,2.67,.91,'white','line',True)
            txt(s,x+.07,y+.12,2.53,.66,label,19.8,'ink',True,'center')
        rect_shape(s,MSO_SHAPE.OVAL,5.25,3.63,2.86,1.92,'blue')
        txt(s,5.46,3.91,2.44,1.23,d['center'],27,'white',True,'center')
        message(s,d['message'],6.27)
    elif k=='accumulation':
        for i,(heading,body) in enumerate(d['steps']):
            x=.83+i*4.25
            rect(s,x,2.51,3.57,2.86,'white','line',True)
            txt(s,x+.21,2.82,3.16,.45,f'{i+1:02d}',19,'blue',True)
            txt(s,x+.21,3.6,3.16,.73,heading,23.5,'ink',True,'center')
            txt(s,x+.21,4.56,3.16,.48,body,17,'muted',False,'center')
            if i<2:arrow(s,x+3.75,3.78,.3,.28)
        message(s,d['message'])
    elif k=='metrics':
        txt(s,.89,2.38,11.55,1.12,d['formula'],48,'blue',True,'center')
        txt(s,.89,3.65,11.55,.51,d['formula_note'],20,'muted',False,'center')
        for i,(label,detail) in enumerate(d['metrics']):
            x=.83+i*4.2
            rect(s,x,4.78,3.6,1.22,'white','line',True)
            txt(s,x+.15,4.97,3.3,.39,label,21,'ink',True,'center')
            txt(s,x+.15,5.5,3.3,.31,detail,15.5,'muted',False,'center')
    elif k=='capacity':
        flow(s,d['steps'],2.53,h=1.15,size=21)
        # Repeated pending tokens illustrate queues without invented throughput data.
        for i in range(5):rect(s,4.52+i*.17,4.39,.11,.28,'amber',radius=True)
        txt(s,3.49,5.04,3.33,.46,'审核排队',20,'amber',True,'center')
        txt(s,7.06,4.54,5.09,.78,'先看谁接手、何时采用\n再安排进入量',22,'ink',True,'center')
        message(s,d['message'])
    elif k=='resources':
        for i,(word,heading,body) in enumerate(d['cards']):
            x=.78+i*4.23
            rect(s,x,2.35,3.6,3.66,'white','line',True)
            txt(s,x+.2,2.64,3.2,1.06,word,58,'blue',True,'center')
            txt(s,x+.2,4.06,3.2,.92,heading,23,'ink',True,'center')
            txt(s,x+.19,5.28,3.22,.4,body,15.2,'muted',False,'center')
    elif k=='implementation':
        for i,(label,detail) in enumerate(d['steps']):
            x=.77+i*3.12
            rect(s,x,2.54,2.75,2.75,'white','line',True)
            txt(s,x+.18,2.81,2.39,.45,f'{i+1:02d}',18,'blue',True)
            txt(s,x+.18,3.49,2.39,.57,label,26,'ink',True,'center')
            txt(s,x+.18,4.33,2.39,.66,detail,17,'muted',False,'center')
            if i<3:arrow(s,x+2.9,3.82,.16,.2)
        message(s,d['message'])
    elif k=='closing':
        txt(s,.82,.58,11.63,.48,'从研究结论，到企业行动',20,'B6C9E4')
        txt(s,.82,1.28,6.0,2.9,'快。',145,'cyan',True)
        txt(s,.95,4.11,11.3,.83,d['subtitle'],35,'white',True)
        for i,label in enumerate(d['speeds']):
            txt(s,1.0+i*2.92,5.4,2.6,.5,label,22,'B6C9E4',True)
        txt(s,.98,6.23,11.42,.44,d['action'],20,'white',True)
    elif k=='appendix_metrics':
        cols=[.78,2.74,7.77];widths=[1.75,4.8,4.74]
        for x,w,label in zip(cols,widths,['类型','一轮记录的起止','一起观察']):
            rect(s,x,2.22,w,.54,'blue')
            txt(s,x+.11,2.3,w-.22,.36,label,17,'white',True)
        for i,row in enumerate(d['rows']):
            yy=2.93+i*.86
            for j,(x,w,value) in enumerate(zip(cols,widths,row)):
                rect(s,x,yy,w,.72,'white' if i%2==0 else 'pale')
                txt(s,x+.13,yy+.1,w-.26,.52,value,18.1 if j==0 else 17.0,'blue' if j==0 else 'ink',j==0)
        txt(s,.94,6.48,11.6,.35,'不同层次各自测量；完整成本、质量和实际采用一起保留。',15,'muted')
    elif k=='sources':
        txt(s,.85,2.15,11.63,.39,'上册依据：第一章执行选择 / 第二章MTL任务 / 第三章流程 / 第四章飞轮 / 第五章实施',16.2,'ink',True)
        for i,(code,name,scope) in enumerate(d['source_rows']):
            yy=2.98+i*.51
            rect(s,.84,yy,11.64,.43,'white' if i%2==0 else 'pale')
            txt(s,.99,yy+.04,.72,.33,code,14.3,'blue',True,url=source_url(code))
            txt(s,1.91,yy+.03,2.05,.35,name,14.7,'ink',True,url=source_url(code))
            txt(s,4.07,yy+.03,8.06,.35,scope,14.7,'muted')
        txt(s,.96,6.27,11.42,.44,'四种快、指标表达与图示，是本演示基于上册的管理归纳。',15,'blue',True)
    else:raise ValueError('Unknown slide kind: '+k)
    return s


def manuscript():
    lines=['# 快：让企业更快完成有效业务循环','',
           '一芥智能研究院｜2026年10月14日｜20页主线＋2页附录｜建议20分钟','',
           '[演示PPT]('+STEM+'.pptx) · [PDF预览]('+STEM+'.pdf) · [制作说明](README.md)','',
           '本讲稿对应2026年10月13日上册v8.0的五章及结语，来源固定为仓库版本`'+DATA['source_revision']+'`。当前上册已调整为七章及结语的v9.0；本演示保留原框架，正文中的章节编号均指v8.0。','',
           '以“快”为解决方案主线；四种快与指标表达为本次演示的管理归纳。案例沿用上册及对应下册的原有来源、阶段与口径。','']
    for d in DATA['slides']:
        lines.extend([f'## {d["number"]:02d}. {d["title"]}','',d['subtitle'],'',d['notes'],'','**依据与出处：**',''])
        for ref in d['refs']:
            src=DATA['sources'][ref]
            lines.append('- ['+src['label']+']('+repository_source_url(ref)+')'+('；阶段：'+src['stage'] if src.get('stage') else '')+'。')
            if 'url' in src:lines.append('  原始出处：[公开来源]('+src['url']+')；定位：'+'；'.join(src['locators'])+'。')
        lines.append('')
    (HERE/'逐页讲稿.md').write_text('\n'.join(lines).rstrip()+'\n',encoding='utf-8')


def validate_pdf(pdf, render_dir):
    import fitz
    from PIL import Image, ImageOps, ImageDraw, ImageFont
    doc=fitz.open(pdf)
    assert len(doc)==len(DATA['slides'])
    compact=lambda text:re.sub(r'\s+','',text)
    missing=[];empty=[];outside=[];duplicate=[]
    scale=1.55
    render_dir.mkdir(parents=True,exist_ok=True)
    for i,page in enumerate(doc):
        text=compact(page.get_text())
        if not text:empty.append(i+1)
        for record in EXPECTED:
            if record['slide']==i+1 and compact(record['text']) not in text:
                missing.append(record)
        for block in page.get_text('dict')['blocks']:
            for row in block.get('lines',[]):
                for span in row.get('spans',[]):
                    x0,y0,x1,y1=span['bbox']
                    if x0 < -.5 or y0 < -.5 or x1 > page.rect.width+.5 or y1 > page.rect.height+.5:
                        outside.append({'slide':i+1,'text':span['text'],'bbox':span['bbox']})
        page.get_pixmap(matrix=fitz.Matrix(scale,scale)).save(render_dir/f'slide-{i+1:02d}.png')
    assert not missing,missing
    assert not empty,empty
    assert not outside,outside
    thumbs=[]
    for i in range(len(doc)):
        im=Image.open(render_dir/f'slide-{i+1:02d}.png').convert('RGB')
        im.thumbnail((384,216))
        tile=Image.new('RGB',(404,244),'#DCE6F3')
        tile.paste(im,((404-im.width)//2,7))
        draw=ImageDraw.Draw(tile);draw.text((14,225),f'{i+1:02d}',fill='#24415E')
        thumbs.append(tile)
    for start in range(0,len(thumbs),12):
        chunk=thumbs[start:start+12];rows=(len(chunk)+3)//4
        sheet=Image.new('RGB',(404*4,244*rows),'#DCE6F3')
        for n,tile in enumerate(chunk):sheet.paste(tile,((n%4)*404,(n//4)*244))
        sheet.save(render_dir/f'overview-{start//12+1}.png')
    return {'pdf_pages':len(doc),'text_boxes_checked':len(EXPECTED),'missing_text_boxes':len(missing),
            'empty_pages':empty,'text_outside_page':outside,'rendered_slides':len(doc)}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--no-pdf',action='store_true',help='Build only editable PPTX and speaker manuscript')
    parser.add_argument('--render-dir',type=Path,default=Path('/tmp/fast-ppt-render'))
    args=parser.parse_args()
    for d in DATA['slides']:render(d)
    assert len(PRS.slides)==22
    for item in SHAPES:
        assert item['x']>=0 and item['y']>=0 and item['x']+item['w']<=W+.001 and item['y']+item['h']<=H+.001,item
    pptx=HERE/(STEM+'.pptx');PRS.save(pptx)
    manuscript()
    manifest={'title':DATA['title'],'date':DATA['date'],'slides':len(PRS.slides),'main_slides':20,
              'font':FONT,'source_revision':DATA['source_revision'],'editable_text_boxes':len(EXPECTED),
              'native_shapes':sum(len(s.shapes) for s in PRS.slides),'raster_images':0,
              'sources':[{'id':key,'path':src['path'],'sha256':hashlib.sha256(source_bytes(src)).hexdigest()} for key,src in DATA['sources'].items()],
              'pptx_sha256':hashlib.sha256(pptx.read_bytes()).hexdigest()}
    if not args.no_pdf:
        subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(HERE),str(pptx)],check=True,capture_output=True,text=True)
        pdf=pptx.with_suffix('.pdf')
        manifest['pdf_sha256']=hashlib.sha256(pdf.read_bytes()).hexdigest()
        manifest['checks']=validate_pdf(pdf,args.render_dir)
    (HERE/'build-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({key:value for key,value in manifest.items() if key!='sources'},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
